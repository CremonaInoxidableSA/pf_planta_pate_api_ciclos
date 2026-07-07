from collections import defaultdict
from datetime import date, datetime, time, timedelta
from bisect import bisect_right
 
from models.ciclo import Ciclo
from models.sensoresAA import SensoresAA
from models.sensores import Sensores
from models.estadoCiclo import EstadoCiclo
from models.receta import Receta
from models.equipo import Equipo
from models.historicoAlarma import HistoricoAlarma
from models.alarmas import Alarmas
from models.alarmasL2 import AlarmasL2
from models.sensoresIO import SensoresIO
 
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
)
from openpyxl.worksheet.table import Table, TableStyleInfo
from io import BytesIO
 
from pathlib import Path
import copy as _copy
import logging
 
logger = logging.getLogger("uvicorn")
 

def generar_reporte_productividad(id_equipo, fecha_inicio, fecha_fin, db):
    """
    Genera el INFORME DE PRODUCTIVIDAD usando obligatoriamente la plantilla
    ``INFORME_DE_PRODUCTIVIDAD.xlsx``.

    Hojas conservadas de la plantilla:
      - PRODUCTIVIDAD
      - LISTA CICLOS
      - LOTE

    Filtros admitidos:
      - 0: todos los equipos.
      - 15: Línea 1 completa.
      - 16: Línea 2 completa.
      - 1..14: equipo puntual. Para LOTE se buscan también los equipos
        complementarios de la misma línea.

    Devuelve:
        BytesIO listo para StreamingResponse.
    """

    # ================================================================
    # 1. CONFIGURACIÓN VERIFICADA CONTRA insert_equipos.sql
    # ================================================================
    IDS_COCINAS_L1 = {1, 2, 3}
    IDS_COCINAS_L2 = {4, 5, 6}
    IDS_ENFRIADORES_L1 = {7, 8, 9, 10}
    IDS_ENFRIADORES_L2 = {11, 12, 13, 14}

    IDS_COCINAS = IDS_COCINAS_L1 | IDS_COCINAS_L2
    IDS_ENFRIADORES = IDS_ENFRIADORES_L1 | IDS_ENFRIADORES_L2
    IDS_LINEA_1 = IDS_COCINAS_L1 | IDS_ENFRIADORES_L1
    IDS_LINEA_2 = IDS_COCINAS_L2 | IDS_ENFRIADORES_L2
    IDS_TODOS = IDS_LINEA_1 | IDS_LINEA_2

    ESTADO_FINALIZADO = "FINALIZADO"
    ESTADO_CANCELADO = "CANCELADO"

    # ================================================================
    # 2. HELPERS GENERALES
    # ================================================================
    def _normalizar_inicio(valor):
        if isinstance(valor, datetime):
            return valor.replace(hour=0, minute=0, second=0, microsecond=0)
        if isinstance(valor, date):
            return datetime.combine(valor, time.min)
        raise TypeError("fecha_inicio debe ser date o datetime")

    def _normalizar_fin(valor):
        if isinstance(valor, datetime):
            return valor.replace(
                hour=23,
                minute=59,
                second=59,
                microsecond=999999,
            )
        if isinstance(valor, date):
            return datetime.combine(valor, time.max)
        raise TypeError("fecha_fin debe ser date o datetime")

    def _tipo_equipo(id_equipo_ciclo):
        if id_equipo_ciclo in IDS_COCINAS:
            return "COCINA"
        if id_equipo_ciclo in IDS_ENFRIADORES:
            return "ENFRIADOR"
        return "DESCONOCIDO"

    def _linea_equipo(id_equipo_ciclo):
        if id_equipo_ciclo in IDS_LINEA_1:
            return "L1"
        if id_equipo_ciclo in IDS_LINEA_2:
            return "L2"
        return "N/A"

    def _segundos_entre(inicio, fin):
        if inicio is None or fin is None:
            return None
        return int((fin - inicio).total_seconds())

    def _duracion_excel(segundos):
        if segundos is None or segundos < 0:
            return None
        return segundos / 86400

    def _copiar_estilo_celda(origen, destino):
        destino.font = _copy.copy(origen.font)
        destino.fill = _copy.copy(origen.fill)
        destino.border = _copy.copy(origen.border)
        destino.alignment = _copy.copy(origen.alignment)
        destino.number_format = origen.number_format
        destino.protection = _copy.copy(origen.protection)

    def _capturar_estilos_fila(hoja, fila, cantidad_columnas):
        return [
            {
                "font": _copy.copy(hoja.cell(fila, columna).font),
                "fill": _copy.copy(hoja.cell(fila, columna).fill),
                "border": _copy.copy(hoja.cell(fila, columna).border),
                "alignment": _copy.copy(
                    hoja.cell(fila, columna).alignment
                ),
                "number_format": hoja.cell(
                    fila,
                    columna,
                ).number_format,
                "protection": _copy.copy(
                    hoja.cell(fila, columna).protection
                ),
            }
            for columna in range(1, cantidad_columnas + 1)
        ]

    def _aplicar_estilo_capturado(celda, estilo):
        celda.font = _copy.copy(estilo["font"])
        celda.fill = _copy.copy(estilo["fill"])
        celda.border = _copy.copy(estilo["border"])
        celda.alignment = _copy.copy(estilo["alignment"])
        celda.number_format = estilo["number_format"]
        celda.protection = _copy.copy(estilo["protection"])

    def _limpiar_rango_datos(
        hoja,
        fila_inicio,
        fila_fin,
        cantidad_columnas,
    ):
        if fila_fin < fila_inicio:
            return
        for fila in range(fila_inicio, fila_fin + 1):
            for columna in range(1, cantidad_columnas + 1):
                hoja.cell(fila, columna).value = None

    def _buscar_tabla_por_rango(hoja, prefijo_rango):
        for tabla in hoja.tables.values():
            if str(tabla.ref).upper().startswith(prefijo_rango.upper()):
                return tabla
        return None

    def _actualizar_rango_tabla(
        hoja,
        tabla,
        nuevo_rango,
        nombre_respaldo,
    ):
        if tabla is None:
            tabla = Table(
                displayName=nombre_respaldo,
                ref=nuevo_rango,
            )
            tabla.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            hoja.add_table(tabla)
        else:
            tabla.ref = nuevo_rango
            if tabla.autoFilter is not None:
                tabla.autoFilter.ref = nuevo_rango

    def _resolver_ruta_plantilla():
        carpeta_servicio = Path(__file__).resolve().parent
        candidatos = (
            carpeta_servicio.parent
            / "data"
            / "INFORME_DE_PRODUCTIVIDAD.xlsx",
            carpeta_servicio
            / "data"
            / "INFORME_DE_PRODUCTIVIDAD.xlsx",
        )
        for candidato in candidatos:
            if candidato.exists():
                return candidato
        raise FileNotFoundError(
            "No se encontró la plantilla INFORME_DE_PRODUCTIVIDAD.xlsx "
            "en la carpeta data del proyecto"
        )

    # ================================================================
    # 3. PERÍODO Y FILTRO DE EQUIPO
    # ================================================================
    fecha_inicio_dt = _normalizar_inicio(fecha_inicio)
    fecha_fin_dt = _normalizar_fin(fecha_fin)

    if fecha_inicio_dt > fecha_fin_dt:
        raise ValueError(
            "La fecha de inicio no puede ser posterior a la fecha de fin"
        )

    if id_equipo == 0:
        ids_lista = set(IDS_TODOS)
        ids_lote = set(IDS_TODOS)
        nombre_filtro_predefinido = "Todos los equipos"
    elif id_equipo == 15:
        ids_lista = set(IDS_LINEA_1)
        ids_lote = set(IDS_LINEA_1)
        nombre_filtro_predefinido = "Todos los equipos - L1"
    elif id_equipo == 16:
        ids_lista = set(IDS_LINEA_2)
        ids_lote = set(IDS_LINEA_2)
        nombre_filtro_predefinido = "Todos los equipos - L2"
    elif id_equipo in IDS_TODOS:
        ids_lista = {id_equipo}
        ids_lote = (
            set(IDS_LINEA_1)
            if id_equipo in IDS_LINEA_1
            else set(IDS_LINEA_2)
        )
        nombre_filtro_predefinido = None
    else:
        raise ValueError(
            "id_equipo inválido. Valores admitidos: 0, 1..14, 15 o 16"
        )

    # ================================================================
    # 4. CONSULTAS AGRUPADAS, SIN N+1
    # ================================================================
    consulta_ciclos = db.query(Ciclo).filter(
        Ciclo.fecha_inicio.isnot(None),
        Ciclo.fecha_fin.isnot(None),
        Ciclo.fecha_fin.between(fecha_inicio_dt, fecha_fin_dt),
    )
    # El filtro 0 representa todos los registros del período, incluso un
    # ciclo legado que haya quedado sin idEquipo. Los demás filtros sí se
    # restringen a los IDs definidos para la línea/equipo.
    if id_equipo != 0:
        consulta_ciclos = consulta_ciclos.filter(
            Ciclo.idEquipo.in_(sorted(ids_lista))
        )

    ciclos_seleccionados = (
        consulta_ciclos
        .order_by(Ciclo.fecha_inicio.asc(), Ciclo.id.asc())
        .all()
    )

    lotes_objetivo = {
        str(ciclo.lote).strip()
        for ciclo in ciclos_seleccionados
        if ciclo.lote is not None and str(ciclo.lote).strip()
    }

    if id_equipo in (0, 15, 16):
        ciclos_para_lotes = list(ciclos_seleccionados)
    else:
        ciclos_complementarios = []
        if lotes_objetivo:
            ciclos_complementarios = (
                db.query(Ciclo)
                .filter(
                    Ciclo.fecha_inicio.isnot(None),
                    Ciclo.fecha_fin.isnot(None),
                    Ciclo.fecha_fin.between(
                        fecha_inicio_dt,
                        fecha_fin_dt,
                    ),
                    Ciclo.idEquipo.in_(sorted(ids_lote)),
                    Ciclo.lote.in_(sorted(lotes_objetivo)),
                )
                .order_by(
                    Ciclo.lote.asc(),
                    Ciclo.fecha_inicio.asc(),
                    Ciclo.id.asc(),
                )
                .all()
            )

        # Los ciclos sin código de lote no pueden buscarse con IN(...),
        # pero deben mantenerse visibles como registros incompletos.
        ciclos_sin_lote_seleccionados = [
            ciclo
            for ciclo in ciclos_seleccionados
            if ciclo.lote is None or not str(ciclo.lote).strip()
        ]

        ciclos_por_id = {
            ciclo.id: ciclo
            for ciclo in (
                ciclos_complementarios
                + ciclos_sin_lote_seleccionados
            )
        }
        ciclos_para_lotes = sorted(
            ciclos_por_id.values(),
            key=lambda ciclo: (
                str(ciclo.lote or ""),
                ciclo.fecha_inicio or datetime.max,
                ciclo.id,
            ),
        )

    ids_equipos_consulta = set(ids_lote)
    ids_equipos_consulta.update(
        ciclo.idEquipo
        for ciclo in ciclos_seleccionados + ciclos_para_lotes
        if ciclo.idEquipo is not None
    )
    equipos = (
        db.query(Equipo)
        .filter(Equipo.id.in_(sorted(ids_equipos_consulta)))
        .all()
        if ids_equipos_consulta
        else []
    )
    equipos_por_id = {equipo.id: equipo for equipo in equipos}

    ids_recetas_consulta = {
        ciclo.idReceta
        for ciclo in ciclos_seleccionados + ciclos_para_lotes
        if ciclo.idReceta is not None
    }
    recetas = (
        db.query(Receta)
        .filter(Receta.id.in_(sorted(ids_recetas_consulta)))
        .all()
        if ids_recetas_consulta
        else []
    )
    recetas_por_id = {receta.id: receta for receta in recetas}

    def _nombre_equipo(id_equipo_ciclo):
        equipo_obj = equipos_por_id.get(id_equipo_ciclo)
        if equipo_obj is None:
            return (
                f"EQUIPO {id_equipo_ciclo}"
                if id_equipo_ciclo is not None
                else "SIN EQUIPO"
            )
        return equipo_obj.nombre

    def _nombre_receta(id_receta):
        receta_obj = recetas_por_id.get(id_receta)
        return receta_obj.nombre if receta_obj is not None else "SIN RECETA"

    if nombre_filtro_predefinido is None:
        nombre_filtro = _nombre_equipo(id_equipo)
    else:
        nombre_filtro = nombre_filtro_predefinido

    # ================================================================
    # 5. LISTA CICLOS
    # ================================================================
    filas_lista_ciclos = []
    cantidad_finalizados = 0
    cantidad_cancelados = 0
    torres_finalizadas = 0
    torres_canceladas = 0
    peso_finalizado = 0.0
    peso_cancelado = 0.0

    for ciclo in ciclos_seleccionados:
        estado = str(ciclo.estadoMaquina or "N/A").strip().upper()
        torres = int(ciclo.cantidadTorres or 0)
        peso = float(ciclo.peso or 0)
        duracion_segundos = _segundos_entre(
            ciclo.fecha_inicio,
            ciclo.fecha_fin,
        )

        if estado == ESTADO_FINALIZADO:
            cantidad_finalizados += 1
            torres_finalizadas += torres
            peso_finalizado += peso
        elif estado == ESTADO_CANCELADO:
            cantidad_cancelados += 1
            torres_canceladas += torres
            peso_cancelado += peso

        filas_lista_ciclos.append([
            ciclo.id,
            _nombre_receta(ciclo.idReceta),
            _nombre_equipo(ciclo.idEquipo),
            _linea_equipo(ciclo.idEquipo),
            str(ciclo.lote).strip()
            if ciclo.lote is not None and str(ciclo.lote).strip()
            else "SIN LOTE",
            peso,
            torres,
            estado,
            ciclo.fecha_inicio,
            ciclo.fecha_fin,
            _duracion_excel(duracion_segundos),
        ])

    # ================================================================
    # 6. EMPAREJAMIENTO DE COCINA Y ENFRIADOR POR LOTE Y LÍNEA
    # ================================================================
    grupos_lote = defaultdict(list)
    for ciclo in ciclos_para_lotes:
        lote_limpio = (
            str(ciclo.lote).strip()
            if ciclo.lote is not None
            else ""
        )
        if lote_limpio:
            clave_lote = lote_limpio
            lote_visible = lote_limpio
        else:
            # No se mezclan todos los ciclos sin lote en un único grupo.
            clave_lote = f"__SIN_LOTE_CICLO_{ciclo.id}"
            lote_visible = "SIN LOTE"

        clave = (clave_lote, _linea_equipo(ciclo.idEquipo))
        grupos_lote[clave].append((ciclo, lote_visible))

    filas_lote = []
    lotes_para_productividad = []

    def _estado_ciclo(ciclo):
        if ciclo is None:
            return None
        return str(ciclo.estadoMaquina or "N/A").strip().upper()

    def _crear_registro_lote(
        lote_visible,
        linea,
        ciclo_cocina,
        ciclo_enfriador,
    ):
        receta_id = (
            ciclo_cocina.idReceta
            if ciclo_cocina is not None
            else (
                ciclo_enfriador.idReceta
                if ciclo_enfriador is not None
                else None
            )
        )
        receta_nombre = _nombre_receta(receta_id)

        nombre_cocina = (
            _nombre_equipo(ciclo_cocina.idEquipo)
            if ciclo_cocina is not None
            else "SIN COCINA"
        )
        nombre_enfriador = (
            _nombre_equipo(ciclo_enfriador.idEquipo)
            if ciclo_enfriador is not None
            else "SIN ENFRIADOR"
        )

        inicio_visible = None
        fin_visible = None
        torres_lote = 0
        peso_lote = 0.0
        tiempo_total_seg = None
        tiempo_muerto_seg = None
        tiempo_util_seg = None
        estado_trazabilidad = "INCOMPLETO"
        cantidad_ciclos = int(ciclo_cocina is not None) + int(
            ciclo_enfriador is not None
        )

        if ciclo_cocina is not None:
            inicio_visible = ciclo_cocina.fecha_inicio
            fin_visible = ciclo_cocina.fecha_fin
            torres_lote = int(ciclo_cocina.cantidadTorres or 0)
            peso_lote = float(ciclo_cocina.peso or 0)

        if ciclo_enfriador is not None:
            if inicio_visible is None:
                inicio_visible = ciclo_enfriador.fecha_inicio
            fin_visible = ciclo_enfriador.fecha_fin
            if ciclo_cocina is None:
                torres_lote = int(ciclo_enfriador.cantidadTorres or 0)
                peso_lote = float(ciclo_enfriador.peso or 0)
            elif peso_lote == 0 and ciclo_enfriador.peso is not None:
                peso_lote = float(ciclo_enfriador.peso)

        if ciclo_cocina is None:
            estado_trazabilidad = "SIN COCINA"
        elif ciclo_enfriador is None:
            estado_trazabilidad = "SIN ENFRIADOR"
        elif (
            ciclo_cocina.fecha_inicio is None
            or ciclo_cocina.fecha_fin is None
            or ciclo_enfriador.fecha_inicio is None
            or ciclo_enfriador.fecha_fin is None
        ):
            estado_trazabilidad = "FECHAS INCOMPLETAS"
        elif (
            _estado_ciclo(ciclo_cocina) == ESTADO_CANCELADO
            or _estado_ciclo(ciclo_enfriador) == ESTADO_CANCELADO
        ):
            estado_trazabilidad = "CICLO CANCELADO"
        elif (
            _estado_ciclo(ciclo_cocina) != ESTADO_FINALIZADO
            or _estado_ciclo(ciclo_enfriador) != ESTADO_FINALIZADO
        ):
            estado_trazabilidad = "ESTADO NO FINALIZADO"
        elif (
            ciclo_cocina.idReceta is not None
            and ciclo_enfriador.idReceta is not None
            and ciclo_cocina.idReceta != ciclo_enfriador.idReceta
        ):
            estado_trazabilidad = "RECETA INCONSISTENTE"
        elif ciclo_enfriador.fecha_inicio < ciclo_cocina.fecha_fin:
            estado_trazabilidad = "INCONSISTENCIA TEMPORAL"
        else:
            tiempo_total_seg = _segundos_entre(
                ciclo_cocina.fecha_inicio,
                ciclo_enfriador.fecha_fin,
            )
            tiempo_muerto_seg = _segundos_entre(
                ciclo_cocina.fecha_fin,
                ciclo_enfriador.fecha_inicio,
            )
            if (
                tiempo_total_seg is not None
                and tiempo_muerto_seg is not None
                and tiempo_total_seg >= tiempo_muerto_seg >= 0
            ):
                tiempo_util_seg = tiempo_total_seg - tiempo_muerto_seg
                estado_trazabilidad = "COMPLETO"
            else:
                estado_trazabilidad = "INCONSISTENCIA TEMPORAL"

        # La plantilla no posee una columna TRAZABILIDAD. Para conservar
        # exactamente su diseño, los estados incompletos se hacen visibles
        # en las columnas COCINA/ENFRIADOR.
        if estado_trazabilidad == "CICLO CANCELADO":
            if _estado_ciclo(ciclo_cocina) == ESTADO_CANCELADO:
                nombre_cocina += " [CANCELADO]"
            if _estado_ciclo(ciclo_enfriador) == ESTADO_CANCELADO:
                nombre_enfriador += " [CANCELADO]"
        elif estado_trazabilidad not in {
            "COMPLETO",
            "SIN COCINA",
            "SIN ENFRIADOR",
        }:
            nombre_enfriador += f" [{estado_trazabilidad}]"

        fila_excel = [
            lote_visible,
            receta_nombre,
            nombre_cocina,
            nombre_enfriador,
            inicio_visible,
            fin_visible,
            torres_lote,
            peso_lote,
            _duracion_excel(tiempo_total_seg),
            _duracion_excel(tiempo_muerto_seg),
            _duracion_excel(tiempo_util_seg),
        ]

        registro_productividad = {
            "lote": lote_visible,
            "linea": linea,
            "id_receta": receta_id,
            "receta": receta_nombre,
            "cantidad_ciclos": cantidad_ciclos,
            "cantidad_torres": torres_lote,
            "peso": peso_lote,
            "tiempo_util_seg": tiempo_util_seg,
            "estado": estado_trazabilidad,
        }
        return fila_excel, registro_productividad

    for (clave_lote, linea), registros_grupo in sorted(
        grupos_lote.items(),
        key=lambda item: (item[0][1], item[0][0]),
    ):
        lote_visible = registros_grupo[0][1]
        ciclos_grupo = [item[0] for item in registros_grupo]

        cocinas = sorted(
            [
                ciclo
                for ciclo in ciclos_grupo
                if ciclo.idEquipo in IDS_COCINAS
            ],
            key=lambda ciclo: (
                ciclo.fecha_inicio or datetime.max,
                ciclo.id,
            ),
        )
        enfriadores_disponibles = sorted(
            [
                ciclo
                for ciclo in ciclos_grupo
                if ciclo.idEquipo in IDS_ENFRIADORES
            ],
            key=lambda ciclo: (
                ciclo.fecha_inicio or datetime.max,
                ciclo.id,
            ),
        )
        ciclos_equipo_desconocido = sorted(
            [
                ciclo
                for ciclo in ciclos_grupo
                if ciclo.idEquipo not in IDS_COCINAS
                and ciclo.idEquipo not in IDS_ENFRIADORES
            ],
            key=lambda ciclo: (
                ciclo.fecha_inicio or datetime.max,
                ciclo.id,
            ),
        )

        # Los equipos desconocidos se muestran como trazabilidad incompleta,
        # pero nunca participan del cálculo de productividad.
        for ciclo_desconocido in ciclos_equipo_desconocido:
            nombre_desconocido = _nombre_equipo(
                ciclo_desconocido.idEquipo
            )
            filas_lote.append([
                lote_visible,
                _nombre_receta(ciclo_desconocido.idReceta),
                "SIN COCINA",
                f"SIN ENFRIADOR [EQUIPO: {nombre_desconocido}]",
                ciclo_desconocido.fecha_inicio,
                ciclo_desconocido.fecha_fin,
                int(ciclo_desconocido.cantidadTorres or 0),
                float(ciclo_desconocido.peso or 0),
                None,
                None,
                None,
            ])
            lotes_para_productividad.append({
                "lote": lote_visible,
                "linea": linea,
                "id_receta": ciclo_desconocido.idReceta,
                "receta": _nombre_receta(
                    ciclo_desconocido.idReceta
                ),
                "cantidad_ciclos": 1,
                "cantidad_torres": int(
                    ciclo_desconocido.cantidadTorres or 0
                ),
                "peso": float(ciclo_desconocido.peso or 0),
                "tiempo_util_seg": None,
                "estado": "EQUIPO DESCONOCIDO",
            })

        # Primera pasada: emparejar solo secuencias temporalmente válidas.
        # Así un enfriador que sirve para una cocina posterior no se consume
        # antes en un emparejamiento inconsistente.
        cocinas_sin_pareja = []
        for cocina in cocinas:
            posteriores = [
                enfriador
                for enfriador in enfriadores_disponibles
                if (
                    enfriador.fecha_inicio is not None
                    and cocina.fecha_fin is not None
                    and enfriador.fecha_inicio >= cocina.fecha_fin
                )
            ]

            if not posteriores:
                cocinas_sin_pareja.append(cocina)
                continue

            enfriador_elegido = min(
                posteriores,
                key=lambda ciclo: (
                    ciclo.fecha_inicio,
                    ciclo.id,
                ),
            )
            enfriadores_disponibles.remove(enfriador_elegido)

            fila_lote, registro_productividad = _crear_registro_lote(
                lote_visible,
                linea,
                cocina,
                enfriador_elegido,
            )
            filas_lote.append(fila_lote)
            lotes_para_productividad.append(registro_productividad)

        # Segunda pasada: si quedaron cocinas y enfriadores del mismo lote,
        # se relacionan para mostrar explícitamente la inconsistencia temporal.
        while cocinas_sin_pareja and enfriadores_disponibles:
            cocina = cocinas_sin_pareja.pop(0)
            enfriador = enfriadores_disponibles.pop(0)
            fila_lote, registro_productividad = _crear_registro_lote(
                lote_visible,
                linea,
                cocina,
                enfriador,
            )
            filas_lote.append(fila_lote)
            lotes_para_productividad.append(registro_productividad)

        # Cocinas sin enfriador disponible.
        for cocina in cocinas_sin_pareja:
            fila_lote, registro_productividad = _crear_registro_lote(
                lote_visible,
                linea,
                cocina,
                None,
            )
            filas_lote.append(fila_lote)
            lotes_para_productividad.append(registro_productividad)

        # Enfriadores sobrantes: no existe cocina disponible para asociar.
        for enfriador in enfriadores_disponibles:
            fila_lote, registro_productividad = _crear_registro_lote(
                lote_visible,
                linea,
                None,
                enfriador,
            )
            filas_lote.append(fila_lote)
            lotes_para_productividad.append(registro_productividad)

        if ciclos_equipo_desconocido:
            logger.warning(
                "Lote %s contiene %s ciclo(s) con equipo desconocido",
                lote_visible,
                len(ciclos_equipo_desconocido),
            )

    filas_lote.sort(
        key=lambda fila: (
            fila[4] or datetime.max,
            str(fila[0]),
        )
    )

    # ================================================================
    # 7. PRODUCTIVIDAD POR RECETA
    # ================================================================
    productividad_por_receta = defaultdict(
        lambda: {
            "cantidad_ciclos": 0,
            "cantidad_torres": 0,
            "tiempo_util_seg": 0,
            "peso_total": 0.0,
        }
    )

    for lote in lotes_para_productividad:
        if lote["estado"] != "COMPLETO":
            continue
        if (
            lote["tiempo_util_seg"] is None
            or lote["tiempo_util_seg"] <= 0
        ):
            continue

        clave_receta = (
            lote["id_receta"],
            lote["receta"],
        )
        acumulado = productividad_por_receta[clave_receta]
        acumulado["cantidad_ciclos"] += lote["cantidad_ciclos"]
        # Torres y peso se suman una sola vez por lote.
        acumulado["cantidad_torres"] += lote["cantidad_torres"]
        acumulado["peso_total"] += lote["peso"]
        acumulado["tiempo_util_seg"] += lote["tiempo_util_seg"]

    filas_productividad = []
    total_ciclos_productivos = 0
    total_torres_productivas = 0
    total_tiempo_util_seg = 0
    total_peso_productivo = 0.0

    for (id_receta, nombre_receta), acumulado in sorted(
        productividad_por_receta.items(),
        key=lambda item: (
            item[0][0] is None,
            item[0][0] or 0,
            item[0][1],
        ),
    ):
        horas_utiles = acumulado["tiempo_util_seg"] / 3600
        kg_hora = (
            acumulado["peso_total"] / horas_utiles
            if horas_utiles > 0
            else 0
        )

        filas_productividad.append([
            id_receta,
            nombre_receta,
            acumulado["cantidad_ciclos"],
            acumulado["cantidad_torres"],
            _duracion_excel(acumulado["tiempo_util_seg"]),
            acumulado["peso_total"],
            round(kg_hora, 2),
        ])

        total_ciclos_productivos += acumulado["cantidad_ciclos"]
        total_torres_productivas += acumulado["cantidad_torres"]
        total_tiempo_util_seg += acumulado["tiempo_util_seg"]
        total_peso_productivo += acumulado["peso_total"]

    horas_utiles_totales = total_tiempo_util_seg / 3600
    kg_hora_total = (
        total_peso_productivo / horas_utiles_totales
        if horas_utiles_totales > 0
        else 0
    )

    # ================================================================
    # 8. CARGA OBLIGATORIA DE LA PLANTILLA
    # ================================================================
    ruta_plantilla = _resolver_ruta_plantilla()
    workbook = load_workbook(ruta_plantilla)

    hojas_requeridas = {
        "PRODUCTIVIDAD",
        "LISTA CICLOS",
        "LOTE",
    }
    hojas_faltantes = hojas_requeridas.difference(workbook.sheetnames)
    if hojas_faltantes:
        workbook.close()
        raise ValueError(
            "La plantilla no contiene las hojas requeridas: "
            + ", ".join(sorted(hojas_faltantes))
        )

    hoja_productividad = workbook["PRODUCTIVIDAD"]
    hoja_lista = workbook["LISTA CICLOS"]
    hoja_lote = workbook["LOTE"]

    # Metadatos superiores. No se alteran títulos, logos, combinaciones,
    # anchos, márgenes ni configuración de impresión de la plantilla.
    for hoja in (hoja_productividad, hoja_lista, hoja_lote):
        hoja["B3"] = fecha_inicio_dt
        hoja["B4"] = fecha_fin_dt
        hoja["B5"] = nombre_filtro
        hoja["B3"].number_format = "dd/mm/yyyy"
        hoja["B4"].number_format = "dd/mm/yyyy"
        hoja.freeze_panes = "A8"

    # ================================================================
    # 9. ESCRITURA DE LISTA CICLOS CONSERVANDO SU DISEÑO
    # ================================================================
    estilos_lista = _capturar_estilos_fila(hoja_lista, 8, 11)
    altura_lista = hoja_lista.row_dimensions[8].height
    fila_fin_limpieza_lista = max(hoja_lista.max_row, 8)
    _limpiar_rango_datos(
        hoja_lista,
        8,
        fila_fin_limpieza_lista,
        11,
    )

    filas_lista_excel = filas_lista_ciclos or [
        ["SIN DATOS"] + [None] * 10
    ]
    for indice, valores in enumerate(filas_lista_excel, start=8):
        for columna, valor in enumerate(valores, start=1):
            celda = hoja_lista.cell(indice, columna, valor)
            _aplicar_estilo_capturado(
                celda,
                estilos_lista[columna - 1],
            )
            if columna in (9, 10) and isinstance(valor, datetime):
                celda.number_format = "yyyy-mm-dd hh:mm:ss"
            elif columna == 11 and valor is not None:
                celda.number_format = "[hh]:mm:ss"
            elif columna == 6 and valor is not None:
                celda.number_format = "0.00"
        if altura_lista is not None:
            hoja_lista.row_dimensions[indice].height = altura_lista

    ultima_fila_lista = 7 + len(filas_lista_excel)
    tabla_lista = _buscar_tabla_por_rango(hoja_lista, "A7:K")
    _actualizar_rango_tabla(
        hoja_lista,
        tabla_lista,
        f"A7:K{ultima_fila_lista}",
        "TablaListaCiclos",
    )

    # Resumen de estados de la plantilla E3:H5.
    hoja_lista["E3"] = "ESTADOS"
    hoja_lista["F3"] = "CANTIDAD CICLOS"
    hoja_lista["G3"] = "CANTIDAD TORRES"
    hoja_lista["H3"] = "TOTAL [KG]"
    hoja_lista["E4"] = "FINALIZADOS"
    hoja_lista["F4"] = cantidad_finalizados
    hoja_lista["G4"] = torres_finalizadas
    hoja_lista["H4"] = peso_finalizado
    hoja_lista["E5"] = "CANCELADOS"
    hoja_lista["F5"] = cantidad_cancelados
    hoja_lista["G5"] = torres_canceladas
    hoja_lista["H5"] = peso_cancelado
    hoja_lista["H4"].number_format = "0.00"
    hoja_lista["H5"].number_format = "0.00"

    # ================================================================
    # 10. ESCRITURA DE LOTE CONSERVANDO SU DISEÑO
    # ================================================================
    estilos_lote = _capturar_estilos_fila(hoja_lote, 8, 11)
    altura_lote = hoja_lote.row_dimensions[8].height
    fila_fin_limpieza_lote = max(hoja_lote.max_row, 8)
    _limpiar_rango_datos(
        hoja_lote,
        8,
        fila_fin_limpieza_lote,
        11,
    )

    filas_lote_excel = filas_lote or [
        ["SIN DATOS"] + [None] * 10
    ]
    for indice, valores in enumerate(filas_lote_excel, start=8):
        for columna, valor in enumerate(valores, start=1):
            celda = hoja_lote.cell(indice, columna, valor)
            _aplicar_estilo_capturado(
                celda,
                estilos_lote[columna - 1],
            )
            if columna in (5, 6) and isinstance(valor, datetime):
                celda.number_format = "yyyy-mm-dd hh:mm:ss"
            elif columna in (9, 10, 11) and valor is not None:
                celda.number_format = "[hh]:mm:ss"
            elif columna == 8 and valor is not None:
                celda.number_format = "0.00"
        if altura_lote is not None:
            hoja_lote.row_dimensions[indice].height = altura_lote

    ultima_fila_lote = 7 + len(filas_lote_excel)
    tabla_lote = _buscar_tabla_por_rango(hoja_lote, "A7:K")
    _actualizar_rango_tabla(
        hoja_lote,
        tabla_lote,
        f"A7:K{ultima_fila_lote}",
        "TablaLotesProductividad",
    )

    # ================================================================
    # 11. ESCRITURA DE PRODUCTIVIDAD Y TOTALES
    # ================================================================
    estilos_productividad = _capturar_estilos_fila(
        hoja_productividad,
        8,
        7,
    )
    estilos_totales = _capturar_estilos_fila(
        hoja_productividad,
        13,
        7,
    )
    altura_productividad = hoja_productividad.row_dimensions[8].height
    altura_totales = hoja_productividad.row_dimensions[13].height
    fila_fin_limpieza_productividad = max(
        hoja_productividad.max_row,
        8,
    )
    _limpiar_rango_datos(
        hoja_productividad,
        8,
        fila_fin_limpieza_productividad,
        7,
    )

    fila_actual = 8
    for valores in filas_productividad:
        for columna, valor in enumerate(valores, start=1):
            celda = hoja_productividad.cell(fila_actual, columna, valor)
            _aplicar_estilo_capturado(
                celda,
                estilos_productividad[columna - 1],
            )
            if columna == 5 and valor is not None:
                celda.number_format = "[hh]:mm:ss"
            elif columna in (6, 7) and valor is not None:
                celda.number_format = "0.00"
        if altura_productividad is not None:
            hoja_productividad.row_dimensions[fila_actual].height = (
                altura_productividad
            )
        fila_actual += 1

    fila_totales = fila_actual
    valores_totales = [
        "TOTALES",
        None,
        total_ciclos_productivos,
        total_torres_productivas,
        _duracion_excel(total_tiempo_util_seg),
        total_peso_productivo,
        round(kg_hora_total, 2),
    ]
    for columna, valor in enumerate(valores_totales, start=1):
        celda = hoja_productividad.cell(fila_totales, columna, valor)
        _aplicar_estilo_capturado(
            celda,
            estilos_totales[columna - 1],
        )
        if columna == 5:
            celda.number_format = "[hh]:mm:ss"
        elif columna in (6, 7):
            celda.number_format = "0.00"
    if altura_totales is not None:
        hoja_productividad.row_dimensions[fila_totales].height = (
            altura_totales
        )

    tabla_productividad = _buscar_tabla_por_rango(
        hoja_productividad,
        "A7:G",
    )
    _actualizar_rango_tabla(
        hoja_productividad,
        tabla_productividad,
        f"A7:G{fila_totales}",
        "ResumenProductividad",
    )

    workbook.active = hoja_productividad

    # ================================================================
    # 12. SALIDA
    # ================================================================
    excel_stream = BytesIO()
    workbook.save(excel_stream)
    workbook.close()
    excel_stream.seek(0)

    logger.info(
        "[generar_reporte_productividad] Reporte generado | "
        "filtro=%s | ciclos=%s | filas_lote=%s | recetas=%s",
        nombre_filtro,
        len(filas_lista_ciclos),
        len(filas_lote),
        len(filas_productividad),
    )

    return excel_stream

